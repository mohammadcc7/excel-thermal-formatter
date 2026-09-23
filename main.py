import os
import sys
import tkinter as tk
from tkinter import filedialog, messagebox
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.utils import get_column_letter

def format_excel_for_thermal(input_path, output_path, remove_empty=True):
    wb = openpyxl.load_workbook(input_path)
    
    for ws in wb.worksheets:
        ws.views.sheetView[0].rightToLeft = True
        
        # إزالة الصفوف فارغة الكمية
        if remove_empty and ws.max_column >= 2:
            rows_to_delete = []
            for r in range(ws.max_row, 1, -1):
                val = ws.cell(row=r, column=2).value
                if val is None or str(val).strip() in ['', '0', 'None']:
                    rows_to_delete.append(r)
            for r in rows_to_delete:
                ws.delete_rows(r)

        max_r = ws.max_row
        max_c = ws.max_column
        
        if max_r == 0 or max_c == 0:
            continue
            
        font_header = Font(name='Arial', size=13, bold=True, color='FFFFFF')
        font_body = Font(name='Arial', size=12, bold=True, color='000000')
        
        fill_header = PatternFill(start_color='000000', end_color='000000', fill_type='solid')
        fill_zebra = PatternFill(start_color='F2F2F2', end_color='F2F2F2', fill_type='solid')
        
        thin_side = Side(style='thin', color='000000')
        border_all = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
        
        for r in range(1, max_r + 1):
            ws.row_dimensions[r].height = 26 if r == 1 else 24
            is_header = (r == 1)
            
            for c in range(1, max_c + 1):
                cell = ws.cell(row=r, column=c)
                
                if isinstance(cell.value, float) and cell.value.is_integer():
                    cell.value = int(cell.value)
                
                cell.font = font_header if is_header else font_body
                cell.border = border_all
                
                if is_header:
                    cell.fill = fill_header
                    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
                else:
                    if r % 2 == 0:
                        cell.fill = fill_zebra
                    
                    if isinstance(cell.value, (int, float)):
                        cell.alignment = Alignment(horizontal='center', vertical='center')
                    else:
                        cell.alignment = Alignment(horizontal='right', vertical='center', wrap_text=True)

        total_units = 38
        if max_c > 1:
            ws.column_dimensions['A'].width = 22
            rem_width = max(8, int((total_units - 22) / (max_c - 1)))
            for c in range(2, max_c + 1):
                ws.column_dimensions[get_column_letter(c)].width = rem_width
        else:
            ws.column_dimensions['A'].width = total_units

        ws.page_margins.left = 0.02
        ws.page_margins.right = 0.02
        ws.page_margins.top = 0.02
        ws.page_margins.bottom = 0.02
        ws.page_margins.header = 0
        ws.page_margins.footer = 0
        
        ws.page_setup.orientation = ws.ORIENTATION_PORTRAIT
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0

    wb.save(output_path)

class App:
    def __init__(self, root):
        self.root = root
        self.root.title("منسق ملفات Excel للطابعة الحرارية (8سم)")
        self.root.geometry("480x320")
        self.root.resizable(False, False)
        
        self.selected_file_path = ""

        # العنوان العلوي
        lbl_title = tk.Label(root, text="محول ومنسق ملفات Excel للطابعة الحرارية", font=("Arial", 13, "bold"), bg="#0056b3", fg="white")
        lbl_title.pack(fill="x", ipady=10)
        
        # خيار الحذف التلقائي
        self.chk_var = tk.BooleanVar(value=True)
        chk = tk.Checkbutton(root, text="حذف الصفوف فارغة الكمية تلقائياً", variable=self.chk_var, font=("Arial", 10))
        chk.pack(pady=10)

        # زر اختيار الملف (أخضر)
        self.btn_select = tk.Button(root, text="اختيار ملف Excel", font=("Arial", 12, "bold"), bg="#28a745", fg="white", command=self.select_file)
        self.btn_select.pack(pady=5, ipadx=20, ipady=5)

        # تسمية لحالة الملف المختار
        self.lbl_status = tk.Label(root, text="لم يتم اختيار أي ملف", font=("Arial", 9), fg="gray")
        self.lbl_status.pack(pady=5)

        # زر تنفيذ المعالجة والتنسيق (أزرق)
        self.btn_process = tk.Button(root, text="تنسيق ومعالجة الملف", font=("Arial", 12, "bold"), bg="#007bff", fg="white", state="disabled", command=self.process_file)
        self.btn_process.pack(pady=10, ipadx=20, ipady=5)

    def select_file(self):
        file_path = filedialog.askopenfilename(filetypes=[("Excel Files", "*.xlsx *.xls")])
        if file_path:
            self.selected_file_path = file_path
            base_name = os.path.basename(file_path)
            self.lbl_status.config(text=f"الملف المختار: {base_name}", fg="green")
            self.btn_process.config(state="normal")

    def process_file(self):
        if not self.selected_file_path:
            return
        
        base_name = os.path.basename(self.selected_file_path)
        save_path = filedialog.asksaveasfilename(defaultextension=".xlsx", initialfile="جاهز_للطباعة_" + base_name, filetypes=[("Excel Files", "*.xlsx")])
        if not save_path:
            return
            
        try:
            format_excel_for_thermal(self.selected_file_path, save_path, remove_empty=self.chk_var.get())
            messagebox.showinfo("نجاح", f"تم تنسيق الملف وحفظه بنجاح:\n{save_path}")
        except Exception as e:
            messagebox.showerror("خطأ", f"حدث خطأ أثناء المعالجة:\n{str(e)}")

if __name__ == "__main__":
    root = tk.Tk()
    app = App(root)
    root.mainloop()
